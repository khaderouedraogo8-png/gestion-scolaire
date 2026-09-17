"""Import bulk élèves (CSV / Excel) — preview + confirm.

Colonnes acceptées (en-têtes normalisés) :
  matricule (optionnel — auto si vide)
  nom (requis)
  prenom (requis)
  sexe (M/F, optionnel)
  date_naissance (YYYY-MM-DD ou JJ/MM/AAAA)
  lieu_naissance
  adresse
  classe (libellé classe) OU id_classe (UUID)

id_annee fourni dans le formulaire (pas dans le fichier).
"""
from __future__ import annotations

import csv
import io
import re
import uuid
from datetime import date, datetime
from typing import Any

from openpyxl import Workbook, load_workbook
from sqlalchemy.orm import Session

from app.models import AnneeScolaire, Classe, Eleve, Etablissement, Inscription
from app.services.tenant import (
    apply_tenant_school,
    assert_same_school,
    get_current_school_id,
    get_or_404_tenant,
    tenant_query,
)
from app.utils.audit_logger import log_audit
from app.utils.errors import abort_api

MAX_IMPORT_ROWS = 2000
MAX_FILE_BYTES = 2 * 1024 * 1024  # 2 Mo

_HEADER_ALIASES = {
    "matricule": "matricule",
    "mat.": "matricule",
    "nom": "nom",
    "prenom": "prenom",
    "prénom": "prenom",
    "sexe": "sexe",
    "date_naissance": "date_naissance",
    "date naissance": "date_naissance",
    "datenais": "date_naissance",
    "lieu_naissance": "lieu_naissance",
    "lieu naissance": "lieu_naissance",
    "adresse": "adresse",
    "classe": "classe",
    "classe_libelle": "classe",
    "libelle_classe": "classe",
    "id_classe": "id_classe",
}

_REQUIRED = ("nom", "prenom")


def _norm_header(raw: str) -> str | None:
    key = re.sub(r"\s+", " ", (raw or "").strip().lower())
    key = key.replace("_", " ")
    # retry with underscore form
    aliases = {k.replace("_", " "): v for k, v in _HEADER_ALIASES.items()}
    aliases.update(_HEADER_ALIASES)
    return aliases.get(key) or aliases.get(key.replace(" ", "_"))


def _parse_date(value: Any) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(text, fmt).date()  # noqa: DTZ007 — date only
        except ValueError:
            continue
    raise ValueError(f"date_naissance invalide : {text}")


def _parse_sexe(value: Any) -> str | None:
    if value is None or str(value).strip() == "":
        return None
    s = str(value).strip().upper()
    if s in ("M", "F"):
        return s
    if s in ("H", "HOMME", "GARCON", "GARÇON", "MASCULIN"):
        return "M"
    if s in ("FEMME", "FILLE", "FEMININ", "FÉMININ"):
        return "F"
    raise ValueError(f"sexe invalide : {value} (attendu M ou F)")


def _cell_str(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _read_tabular(file_storage) -> list[dict[str, Any]]:
    """Lit CSV ou XLSX → liste de dicts bruts (clés normalisées)."""
    filename = (getattr(file_storage, "filename", "") or "").lower()
    raw = file_storage.read()
    if hasattr(file_storage, "seek"):
        file_storage.seek(0)
    if not raw:
        abort_api(400, "IMPORT_EMPTY", "Fichier vide.")
    if len(raw) > MAX_FILE_BYTES:
        abort_api(400, "IMPORT_TOO_LARGE", f"Fichier trop volumineux (max {MAX_FILE_BYTES // (1024 * 1024)} Mo).")

    rows_raw: list[list[Any]]
    if filename.endswith(".csv"):
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode("latin-1")
        sample = text[:2048]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=";,\t")
        except csv.Error:
            dialect = csv.excel
            dialect.delimiter = ";" if sample.count(";") >= sample.count(",") else ","
        reader = csv.reader(io.StringIO(text), dialect)
        rows_raw = list(reader)
    elif filename.endswith((".xlsx", ".xlsm")):
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        ws = wb.active
        rows_raw = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()
    else:
        abort_api(
            400,
            "IMPORT_BAD_EXTENSION",
            "Extension non supportée. Utilisez .xlsx ou .csv.",
        )

    if not rows_raw:
        abort_api(400, "IMPORT_EMPTY", "Fichier sans lignes.")

    # Skip leading empty rows
    header_idx = 0
    while header_idx < len(rows_raw) and all(
        c is None or str(c).strip() == "" for c in rows_raw[header_idx]
    ):
        header_idx += 1
    if header_idx >= len(rows_raw):
        abort_api(400, "IMPORT_EMPTY", "Fichier sans en-têtes.")

    headers_raw = rows_raw[header_idx]
    headers: list[str | None] = [_norm_header(_cell_str(h)) for h in headers_raw]
    if not any(headers):
        abort_api(400, "IMPORT_BAD_HEADERS", "Aucune colonne reconnue.")

    mapped = [h for h in headers if h]
    for req in _REQUIRED:
        if req not in mapped:
            abort_api(
                400,
                "IMPORT_MISSING_COLUMN",
                f"Colonne obligatoire manquante : {req}.",
                details={"required": list(_REQUIRED), "found": mapped},
            )

    data_rows: list[dict[str, Any]] = []
    for offset, row in enumerate(rows_raw[header_idx + 1 :], start=header_idx + 2):
        if all(c is None or str(c).strip() == "" for c in row):
            continue
        item: dict[str, Any] = {"_line": offset}
        for i, key in enumerate(headers):
            if not key:
                continue
            val = row[i] if i < len(row) else None
            item[key] = val
        data_rows.append(item)

    if len(data_rows) > MAX_IMPORT_ROWS:
        abort_api(
            400,
            "IMPORT_TOO_MANY_ROWS",
            f"Trop de lignes (max {MAX_IMPORT_ROWS}).",
            details={"count": len(data_rows)},
        )
    if not data_rows:
        abort_api(400, "IMPORT_EMPTY", "Aucune ligne de données.")
    return data_rows


def _classe_lookup(db: Session) -> dict[str, Classe]:
    classes = tenant_query(Classe).all()
    by_label = {}
    for c in classes:
        by_label[c.libelle.strip().lower()] = c
        by_label[str(c.id)] = c
    return by_label


def _resolve_classe(
    db: Session,
    *,
    id_classe: str | None,
    classe_libelle: str | None,
    lookup: dict[str, Classe],
    default_classe_id: uuid.UUID | None,
) -> Classe:
    if id_classe:
        cid = str(id_classe).strip()
        try:
            uid = uuid.UUID(cid)
            c = tenant_query(Classe).filter(Classe.id == uid).first()
            if c:
                return c
            raise ValueError(f"Classe introuvable : {cid}")
        except ValueError as exc:
            if "Classe introuvable" in str(exc):
                raise
            c = lookup.get(cid.lower())
            if c:
                return c
            raise ValueError(f"Classe introuvable : {cid}") from exc
    if classe_libelle:
        key = classe_libelle.strip().lower()
        c = lookup.get(key)
        if not c:
            raise ValueError(f"Classe introuvable : {classe_libelle}")
        return c
    if default_classe_id:
        c = tenant_query(Classe).filter(Classe.id == default_classe_id).first()
        if not c:
            raise ValueError("Classe par défaut introuvable.")
        return c
    raise ValueError("Classe manquante (colonne classe ou id_classe).")


def _generer_matricule(db: Session, id_annee: uuid.UUID, seq: int) -> str:
    school_id = get_current_school_id()
    etab = db.query(Etablissement).filter(Etablissement.school_id == school_id).first()
    annee = get_or_404_tenant(AnneeScolaire, id_annee)
    format_str = (etab.format_matricule if etab else None) or "{ANNEE}M-{SEQ}"
    annee_court = annee.libelle[:4] if annee else str(date.today().year)
    return format_str.replace("{ANNEE}", annee_court).replace("{SEQ}", str(seq).zfill(3))


def validate_import_rows(
    db: Session,
    raw_rows: list[dict[str, Any]],
    *,
    id_annee: uuid.UUID,
    default_id_classe: uuid.UUID | None = None,
) -> dict[str, Any]:
    get_or_404_tenant(AnneeScolaire, id_annee)
    if default_id_classe:
        get_or_404_tenant(Classe, default_id_classe)

    lookup = _classe_lookup(db)
    existing_matricules = {
        e.matricule.strip().lower()
        for e in tenant_query(Eleve).with_entities(Eleve.matricule).all()
    }
    existing_identity = {
        (
            (e.nom or "").strip().lower(),
            (e.prenom or "").strip().lower(),
            e.date_naissance.isoformat() if e.date_naissance else "",
        ): e.matricule
        for e in tenant_query(Eleve).all()
    }
    seen_in_file: set[str] = set()
    seen_identity_file: set[tuple[str, str, str]] = set()
    doublons_suspects: list[dict[str, Any]] = []
    base_count = tenant_query(Eleve).count()
    auto_seq = base_count

    results: list[dict[str, Any]] = []
    ok_count = 0
    err_count = 0

    for raw in raw_rows:
        line = raw.get("_line")
        errors: list[str] = []
        warnings: list[str] = []
        data: dict[str, Any] = {}

        nom = _cell_str(raw.get("nom"))
        prenom = _cell_str(raw.get("prenom"))
        if not nom:
            errors.append("nom obligatoire")
        if not prenom:
            errors.append("prenom obligatoire")
        data["nom"] = nom
        data["prenom"] = prenom

        try:
            data["sexe"] = _parse_sexe(raw.get("sexe"))
        except ValueError as exc:
            errors.append(str(exc))

        try:
            data["date_naissance"] = _parse_date(raw.get("date_naissance"))
            if data["date_naissance"]:
                data["date_naissance"] = data["date_naissance"].isoformat()
        except ValueError as exc:
            errors.append(str(exc))

        data["lieu_naissance"] = _cell_str(raw.get("lieu_naissance")) or None
        data["adresse"] = _cell_str(raw.get("adresse")) or None

        identity_key = (
            nom.strip().lower(),
            prenom.strip().lower(),
            data.get("date_naissance") or "",
        )
        if nom and prenom and identity_key[2]:
            if identity_key in existing_identity:
                msg = (
                    f"doublon identité suspect (DB matricule "
                    f"{existing_identity[identity_key]}) : {prenom} {nom} "
                    f"né(e) {identity_key[2]}"
                )
                errors.append(msg)
                doublons_suspects.append({
                    "line": line,
                    "source": "database",
                    "nom": nom,
                    "prenom": prenom,
                    "date_naissance": identity_key[2],
                    "matricule_existant": existing_identity[identity_key],
                })
            elif identity_key in seen_identity_file:
                msg = f"doublon identité dans le fichier : {prenom} {nom} né(e) {identity_key[2]}"
                errors.append(msg)
                doublons_suspects.append({
                    "line": line,
                    "source": "file",
                    "nom": nom,
                    "prenom": prenom,
                    "date_naissance": identity_key[2],
                })
            else:
                seen_identity_file.add(identity_key)
        elif nom and prenom and not identity_key[2]:
            warnings.append(
                "date_naissance absente — dédup identité (nom+prénom+DOB) non applicable"
            )

        matricule = _cell_str(raw.get("matricule"))
        if matricule:
            key = matricule.lower()
            if key in existing_matricules:
                errors.append(f"matricule déjà existant : {matricule}")
            elif key in seen_in_file:
                errors.append(f"matricule dupliqué dans le fichier : {matricule}")
            else:
                seen_in_file.add(key)
            data["matricule"] = matricule
            data["matricule_auto"] = False
        else:
            auto_seq += 1
            data["matricule"] = _generer_matricule(db, id_annee, auto_seq)
            data["matricule_auto"] = True
            while data["matricule"].lower() in existing_matricules or data["matricule"].lower() in seen_in_file:
                auto_seq += 1
                data["matricule"] = _generer_matricule(db, id_annee, auto_seq)
            seen_in_file.add(data["matricule"].lower())

        try:
            classe = _resolve_classe(
                db,
                id_classe=_cell_str(raw.get("id_classe")) or None,
                classe_libelle=_cell_str(raw.get("classe")) or None,
                lookup=lookup,
                default_classe_id=default_id_classe,
            )
            data["id_classe"] = str(classe.id)
            data["classe_libelle"] = classe.libelle
        except ValueError as exc:
            errors.append(str(exc))

        data["id_annee"] = str(id_annee)
        status = "ok" if not errors else "error"
        if status == "ok":
            ok_count += 1
        else:
            err_count += 1
        results.append({
            "line": line,
            "status": status,
            "errors": errors,
            "warnings": warnings,
            "data": data,
        })

    return {
        "items": results,
        "summary": {
            "total": len(results),
            "ok": ok_count,
            "errors": err_count,
            "doublons_suspects": len(doublons_suspects),
        },
        "doublons_suspects": doublons_suspects,
        "id_annee": str(id_annee),
    }


def preview_eleves_import(
    db: Session,
    file_storage,
    *,
    id_annee: uuid.UUID,
    default_id_classe: uuid.UUID | None = None,
) -> dict[str, Any]:
    raw_rows = _read_tabular(file_storage)
    return validate_import_rows(
        db, raw_rows, id_annee=id_annee, default_id_classe=default_id_classe
    )


def confirm_eleves_import(
    db: Session,
    *,
    id_annee: uuid.UUID,
    rows: list[dict[str, Any]],
    user_id: uuid.UUID | None,
) -> dict[str, Any]:
    """Ré-valide puis importe en transaction unique."""
    if not rows:
        abort_api(400, "IMPORT_EMPTY", "Aucune ligne à importer.")
    if len(rows) > MAX_IMPORT_ROWS:
        abort_api(400, "IMPORT_TOO_MANY_ROWS", f"Trop de lignes (max {MAX_IMPORT_ROWS}).")

    # Rebuild raw-like for revalidation
    raw_rows = []
    for i, r in enumerate(rows, start=1):
        raw_rows.append(
            {
                "_line": r.get("line") or i,
                "matricule": r.get("matricule") if not r.get("matricule_auto") else "",
                "nom": r.get("nom"),
                "prenom": r.get("prenom"),
                "sexe": r.get("sexe"),
                "date_naissance": r.get("date_naissance"),
                "lieu_naissance": r.get("lieu_naissance"),
                "adresse": r.get("adresse"),
                "id_classe": r.get("id_classe"),
                "classe": r.get("classe_libelle") or r.get("classe"),
            }
        )

    preview = validate_import_rows(db, raw_rows, id_annee=id_annee)
    if preview["summary"]["errors"]:
        abort_api(
            400,
            "IMPORT_VALIDATION_FAILED",
            "Des lignes sont encore invalides — corrigez puis réessayez.",
            details={"summary": preview["summary"], "items": preview["items"]},
        )

    annee = get_or_404_tenant(AnneeScolaire, id_annee)
    created: list[dict[str, Any]] = []
    try:
        for item in preview["items"]:
            data = item["data"]
            classe = get_or_404_tenant(Classe, uuid.UUID(data["id_classe"]))
            assert_same_school(classe, annee)

            # Prefer provided matricule (including previously auto ones from preview)
            matricule = data["matricule"]
            exists = (
                tenant_query(Eleve)
                .filter(Eleve.matricule == matricule)
                .first()
            )
            if exists:
                abort_api(
                    409,
                    "IMPORT_MATRICULE_CONFLICT",
                    f"Conflit matricule pendant l'import : {matricule}",
                )

            dn = data.get("date_naissance")
            if isinstance(dn, str) and dn:
                dn = date.fromisoformat(dn)
            else:
                dn = None

            eleve = Eleve(
                id=uuid.uuid4(),
                matricule=matricule,
                nom=data["nom"],
                prenom=data["prenom"],
                sexe=data.get("sexe"),
                date_naissance=dn,
                lieu_naissance=data.get("lieu_naissance"),
                adresse=data.get("adresse"),
            )
            apply_tenant_school(eleve)
            db.add(eleve)

            inscription = Inscription(
                id=uuid.uuid4(),
                id_eleve=eleve.id,
                id_classe=classe.id,
                id_annee=annee.id,
                statut="inscrit",
            )
            apply_tenant_school(inscription)
            assert_same_school(eleve, classe, annee, inscription)
            db.add(inscription)
            created.append(
                {
                    "id": str(eleve.id),
                    "matricule": eleve.matricule,
                    "nom": eleve.nom,
                    "prenom": eleve.prenom,
                    "id_classe": str(classe.id),
                }
            )

        db.commit()
    except Exception:
        db.rollback()
        raise

    if user_id:
        log_audit("IMPORT_ELEVES", user_id, "eleve", None)
    return {"imported": len(created), "items": created}


def build_import_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Eleves"
    ws.append(
        [
            "matricule",
            "nom",
            "prenom",
            "sexe",
            "date_naissance",
            "lieu_naissance",
            "adresse",
            "classe",
        ]
    )
    ws.append(
        [
            "",
            "OUEDRAOGO",
            "Awa",
            "F",
            "2012-05-14",
            "Ouagadougou",
            "Secteur 15",
            "6ème A",
        ]
    )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
